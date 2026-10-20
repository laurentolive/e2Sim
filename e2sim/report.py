"""Excel traceability report of a pytest session (spec §15).

Same layout as the e2Unit report (sheets "Summary" and "Tests", same columns and colours), so that the unit
(e2Unit) and system (e2Sim) reports merge into one requirement -> test -> verdict matrix; plus a
"Requirements" sheet. Shared format, no code dependency between the repositories.
"""

from __future__ import annotations

import datetime
from dataclasses import dataclass, field
from pathlib import Path

PASS, FAIL, ERROR, SKIP = "PASS", "FAIL", "ERROR", "SKIP"
_FILL = {PASS: "C6EFCE", FAIL: "FFC7CE", ERROR: "FFC7CE", SKIP: "EDEDED"}
_RANK = {PASS: 0, SKIP: 1, FAIL: 2, ERROR: 3}


@dataclass
class TestRecord:
    id: str
    title: str
    function: str
    requirements: list[str]
    status: str = PASS
    message: str = ""
    duration: float = 0.0


@dataclass
class SessionReport:
    name: str
    started: str = field(default_factory=lambda: datetime.datetime.now().isoformat(timespec="seconds"))
    license: str = ""
    toolchain: str = ""
    tests: list[TestRecord] = field(default_factory=list)

    def count(self, status: str) -> int:
        return sum(1 for t in self.tests if t.status == status)

    @property
    def status(self) -> str:
        if self.count(ERROR):
            return ERROR
        if self.count(FAIL):
            return FAIL
        return PASS

    def requirements(self) -> dict[str, list[TestRecord]]:
        out: dict[str, list[TestRecord]] = {}
        for t in self.tests:
            for r in t.requirements or ["(none)"]:
                out.setdefault(r, []).append(t)
        return dict(sorted(out.items()))


def write_xlsx(rep: SessionReport, path: str | Path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as e:
        raise RuntimeError("openpyxl is required for the Excel report: pip install openpyxl") from e
    from . import __version__

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")
    fill = {k: PatternFill("solid", fgColor=v) for k, v in _FILL.items()}
    duration = round(sum(t.duration for t in rep.tests), 3)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for row in (("Suite", rep.name), ("Status", rep.status), ("Started", rep.started), ("Duration (s)", duration),
                ("Passed", rep.count(PASS)), ("Failed", rep.count(FAIL)), ("Errors", rep.count(ERROR)),
                ("Skipped", rep.count(SKIP)), ("e2Sim", __version__), ("License", rep.license),
                ("Toolchain", rep.toolchain)):
        ws.append(row)
    for c in ws["A"]:
        c.font = Font(bold=True)
    ws["B2"].fill = fill[rep.status]
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 90

    wt = wb.create_sheet("Tests")
    wt.append(["Test Id", "Requirement", "Title", "Function", "Test status", "Step", "Check", "Step status",
               "Result", "Duration (s)"])
    for c in wt[1]:
        c.font, c.fill = head_font, head_fill
    for t in rep.tests:
        wt.append([t.id, ", ".join(t.requirements), t.title, t.function, t.status, f"{t.id}.1", "pytest scenario",
                   t.status, t.message, round(t.duration, 3)])
        row = wt.max_row
        wt.cell(row, 5).fill = fill[t.status]
        wt.cell(row, 8).fill = fill[t.status]
        wt.cell(row, 9).alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCDEFGHIJ", (34, 22, 50, 40, 11, 14, 16, 11, 90, 12)):
        wt.column_dimensions[col].width = width
    wt.freeze_panes = "A2"
    wt.auto_filter.ref = wt.dimensions

    wr = wb.create_sheet("Requirements")
    wr.append(["Requirement", "Status", "Tests", "Verdicts"])
    for c in wr[1]:
        c.font, c.fill = head_font, head_fill
    for req, tests in rep.requirements().items():
        worst = max((t.status for t in tests), key=_RANK.get)
        wr.append([req, worst, "\n".join(t.id for t in tests), "\n".join(t.status for t in tests)])
        row = wr.max_row
        wr.cell(row, 2).fill = fill[worst]
        for col in (3, 4):
            wr.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCD", (24, 11, 50, 11)):
        wr.column_dimensions[col].width = width
    wr.freeze_panes = "A2"

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
