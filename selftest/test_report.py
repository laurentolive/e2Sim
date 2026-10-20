"""Excel report: statuses and requirement matrix."""

import openpyxl

from e2sim.report import ERROR, FAIL, PASS, SKIP, SessionReport, TestRecord, write_xlsx


def test_report_statuses(tmp_path):
    rep = SessionReport(name="demo", license="Demonstration")
    rep.tests += [TestRecord("t1", "ok", "f::t1", ["REQ-1"]), TestRecord("t2", "ko", "f::t2", ["REQ-1", "REQ-2"], FAIL, "boom"),
                  TestRecord("t3", "skip", "f::t3", [], SKIP), TestRecord("t4", "err", "f::t4", ["REQ-2"], ERROR)]
    assert rep.status == ERROR
    write_xlsx(rep, tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "r.xlsx")
    assert wb["Summary"]["B2"].value == ERROR
    reqs = {r[0]: r[1] for r in wb["Requirements"].iter_rows(min_row=2, values_only=True)}
    assert reqs == {"(none)": SKIP, "REQ-1": FAIL, "REQ-2": ERROR}
    rows = list(wb["Tests"].iter_rows(min_row=2, values_only=True))
    assert [r[4] for r in rows] == [PASS, FAIL, SKIP, ERROR] and rows[1][8] == "boom"
